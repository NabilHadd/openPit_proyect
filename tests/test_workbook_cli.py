from contextlib import redirect_stdout
from copy import copy
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, PatternFill

from main import execute, parser, reference_report
from src.config import PROJECT_ROOT, ProblemParameters
from src.data.dataset import file_hash, load_mine
from src.data.transforms import grid9_z_up, load_precedences, load_selections
from src.evaluation.metrics import Extraction, Plan
from src.evaluation.benchmark import maximum_closure
from src.evaluation.reporter import export_checked_plan, load_plan, plan_payload, write_json


def synthetic_workbook(path):
    """Fabricated data, NOT the client's workbook or its economic reference."""
    workbook = Workbook()
    table = workbook.active
    table.title = "105 Bloques"
    table.append(["ID", "X", "Y", "Z", "Valor"])
    for z in range(1, 4):
        for y in range(1, 6):
            for x in range(1, 8):
                block_id = x + (y - 1) * 7 + (z - 1) * 35
                table.append([block_id, x, y, z, 2])
    for name, count, fill in (
        ("X-Y View", 53, PatternFill("solid", fgColor="FFFFFF00")),
        ("Resultado", 42, PatternFill("solid", fgColor=Color(theme=8))),
    ):
        sheet = workbook.create_sheet(name)
        sheet["K3"].fill = fill
        for z, label_row, header_row in ((3, 1, 2), (2, 10, 10), (1, 18, 18)):
            sheet.cell(label_row, 1, f"Z={z}")
            for x in range(1, 8):
                sheet.cell(header_row, x + 2, x)
            for offset, y in enumerate(range(5, 0, -1), 1):
                sheet.cell(header_row + offset, 2, y)
                for x in range(1, 8):
                    block_id = x + (y - 1) * 7 + (z - 1) * 35
                    cell = sheet.cell(header_row + offset, x + 2, block_id)
                    if block_id <= count:
                        cell.fill = fill
        sheet["K6" if name == "X-Y View" else "L13"] = count if name == "X-Y View" else count * 2
    workbook.save(path)


class WorkbookTests(unittest.TestCase):
    def test_selections_traceability_and_correspondence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            before = file_hash(path)
            mine = load_mine(path)
            selections = load_selections(path, mine)
            self.assertEqual((len(mine.blocks), len(selections.candidates.ids), len(selections.reference.ids)), (105, 53, 42))
            report, _ = reference_report(mine, selections, ProblemParameters(), None)
            self.assertEqual(report["recalculated_raw_value"], 84)
            self.assertTrue(report["economic_sum_matches"])
            self.assertIsNone(report["evaluation"]["feasible"])
            self.assertEqual(file_hash(path), before)
            workbook = load_workbook(path)
            workbook["Resultado"]["C3"] = 999
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "mismo archivo"):
                load_selections(path, mine)
            with self.assertRaisesRegex(ValueError, "no corresponden"):
                load_selections(path, load_mine(path))

    def test_explicit_precedence_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            mine = load_mine(path)
            rules = Path(directory) / "rules.json"
            write_json(rules, {"source": "Synthetic: independent blocks", "predecessors": {str(i): [] for i in mine.by_id}})
            precedences = load_precedences(rules, mine)
            self.assertFalse(precedences.provisional)
            self.assertEqual(len(precedences.predecessors), 105)
            write_json(rules, {"source": "test", "predecessors": {"1": [999]}})
            with self.assertRaises(ValueError):
                load_precedences(rules, mine)

    def test_reference_cannot_influence_training(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            args = parser().parse_args(["train", "--data", str(path), "--output", directory,
                                        "--precedence", "grid9-z-up", "--episodes", "6", "--seeds", "9"])
            with redirect_stdout(io.StringIO()):
                first = execute(args)
                workbook = load_workbook(path)
                workbook["Resultado"]["A21"] = "Unrelated note"
                workbook["Resultado"]["L13"] = 99999
                workbook["Resultado"]["C3"].fill = copy(workbook["Resultado"]["K3"].fill)
                workbook["X-Y View"]["K6"] = 999
                workbook.save(path)
                second = execute(args)
            for name in ("training.json", "q-table.json"):
                self.assertEqual((first / "seed-9" / name).read_bytes(), (second / "seed-9" / name).read_bytes())
            report = json.loads((first / "report.json").read_text())
            self.assertEqual(report["data"]["block_count"], 105)
            self.assertFalse(report["validation_metadata"]["used_by_agent"])
            self.assertIn("source_sha256", report["code"])
            self.assertEqual(report["effective_config"]["learning"]["episodes"], 6)
            # Fabricated reference selects lower blocks without their upper blocks.
            self.assertFalse(report["comparison"]["reference_feasible"])
            self.assertFalse(report["comparison"]["comparable"])
            self.assertIsNone(report["comparison"]["best_episode_target_hits"])

    def test_internal_plan_roundtrip_and_export_gate(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "fabricated.xlsx"
            synthetic_workbook(path)
            mine = load_mine(path)
            precedences = grid9_z_up(mine)
            plan = Plan((Extraction(105),))
            write_json(root / "input.json", plan_payload(plan))
            self.assertEqual(load_plan(root / "input.json"), plan)
            self.assertIsNone(export_checked_plan(root, mine, plan, ProblemParameters(), None))
            self.assertEqual(export_checked_plan(root, mine, plan, ProblemParameters(), precedences), "conditional-plan.json")
            self.assertIsNone(export_checked_plan(root, mine, Plan((Extraction(1),)), ProblemParameters(), precedences))

    def test_spatial_reference_and_benchmark_are_isolated(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            args = parser().parse_args(["train", "--data", str(path), "--output", directory,
                                        "--precedence", "grid9-z-up", "--model", "spatial-cone",
                                        "--episodes", "6", "--max-steps", "30", "--seeds", "9"])
            with redirect_stdout(io.StringIO()):
                first = execute(args)
                workbook = load_workbook(path)
                workbook["Resultado"]["L13"] = 99999
                workbook["Resultado"]["C3"].fill = copy(workbook["Resultado"]["K3"].fill)
                workbook["X-Y View"]["K6"] = 999
                workbook.save(path)
                args.exact_benchmark = True
                second = execute(args)
            for name in ("training.json", "q-table.json"):
                self.assertEqual((first / "seed-9" / name).read_bytes(), (second / "seed-9" / name).read_bytes())
            report = json.loads((second / "report.json").read_text())
            self.assertEqual(report["effective_config"]["learning"], {"episodes": 6, "alpha": 1.0, "gamma": 0.7, "epsilon": 0.3})
            self.assertFalse(report["benchmark_comparison"]["used_by_agent"])
            self.assertFalse(report["comparison"]["reference_matches_declaration"])
            entry = report["experiments"][0]
            selected_file = second / "seed-9" / "selected" / entry["selected_export"]
            exported = json.loads(selected_file.read_text())
            self.assertEqual(plan_payload(load_plan(selected_file))["extractions"], entry["selected_path"]["plan"]["extractions"])
            self.assertEqual(exported["evaluation"], entry["selected_path"]["evaluation"])
            self.assertTrue(exported["evaluation"]["feasible"])
            self.assertEqual(len(entry["selected_path"]["actions"]), entry["selected_path"]["steps"])
            self.assertIn("B:", report["domain_confirmations"]["reward"])
            args.surface = "min-z"
            with self.assertRaisesRegex(ValueError, "contradice"):
                execute(args)

    def test_benchmark_cli_without_visual_reference_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            workbook = load_workbook(path)
            del workbook["Resultado"]
            del workbook["X-Y View"]
            workbook["105 Bloques"]["E2"] = 1.25
            workbook.save(path)
            args = parser().parse_args(["benchmark", "--data", str(path), "--output", directory, "--precedence", "grid9-z-up"])
            with redirect_stdout(io.StringIO()):
                directory = execute(args)
            report = json.loads((directory / "report.json").read_text())
            self.assertEqual(report["evaluation"]["npv"], 209.25)
            self.assertNotIn("reference", report)
            self.assertEqual(report["exact_benchmark"]["exact_value"], "837/4")

    def test_cli_outside_repository_and_missing_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fabricated.xlsx"
            synthetic_workbook(path)
            command = [sys.executable, str(PROJECT_ROOT / "main.py"), "reference", "--data", str(path), "--output", directory]
            result = subprocess.run(command, cwd=directory, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("VAN recalculado: 84", result.stdout)
            result = subprocess.run([sys.executable, str(PROJECT_ROOT / "main.py"), "train", "--data", str(path)],
                                    cwd=directory, capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertIn("precedencias explícitas", result.stderr)

    def test_wrong_json_types_are_controlled_cli_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "fabricated.xlsx"
            synthetic_workbook(path)
            bad = root / "bad.json"
            for payload in ([], None):
                write_json(bad, payload)
                commands = (
                    ["evaluate-plan", "--plan", str(bad)],
                    ["train", "--precedences-file", str(bad)],
                )
                for command in commands:
                    with self.subTest(payload=payload, command=command[0]):
                        result = subprocess.run(
                            [sys.executable, str(PROJECT_ROOT / "main.py"), *command,
                             "--data", str(path), "--output", directory],
                            cwd=directory, capture_output=True, text=True,
                        )
                        self.assertEqual(result.returncode, 2, result.stderr)
                        self.assertIn("objeto", result.stderr)
                        self.assertNotIn("Traceback", result.stderr)


@unittest.skipUnless(os.environ.get("OPENPIT_TEST_XLSX"), "Excel privado no solicitado; usar OPENPIT_TEST_XLSX con ruta explícita")
class PrivateWorkbookIntegration(unittest.TestCase):
    def test_real_counts_economics_and_conditional_feasibility(self):
        path = Path(os.environ["OPENPIT_TEST_XLSX"]).resolve(strict=True)
        before = file_hash(path)
        mine = load_mine(path)
        selections = load_selections(path, mine)
        report, _ = reference_report(mine, selections, ProblemParameters(), grid9_z_up(mine))
        self.assertEqual((len(mine.blocks), len(selections.candidates.ids), len(selections.reference.ids)), (105, 53, 42))
        self.assertEqual(report["recalculated_raw_value"], 336000)
        self.assertEqual(report["declared_value"], 336000)
        self.assertTrue(report["economic_sum_matches"])
        self.assertTrue(report["evaluation"]["feasible"])
        self.assertEqual(report["evaluation"]["precedence_status"], "conditional")
        self.assertTrue(report["candidate_geometry_diagnostic"]["matches_excel_candidates"])
        benchmark = maximum_closure(mine, grid9_z_up(mine))
        self.assertEqual(benchmark.exact_value, "336000")
        self.assertEqual(len(benchmark.plan.extractions), 42)
        self.assertEqual(file_hash(path), before)


if __name__ == "__main__":
    unittest.main()
