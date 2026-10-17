import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Color, PatternFill

from src.config import LearningParameters, ProblemParameters
from src.data.dataset import Block, Mine, load_mine
from src.data.transforms import Precedences, color_rgb, fill_signature, grid9_z_up, order_selection
from src.evaluation.metrics import Extraction, Plan, evaluate_plan


def small_mine():
    return Mine((Block(10, 1, 1, 2, -4, tonnes=2), Block(20, 1, 1, 1, 15, tonnes=3)))


class DataTests(unittest.TestCase):
    def test_parameter_validation(self):
        invalid_problems = ({"periods": 0}, {"discount_rate": float("nan")},
                            {"discount_rate": -0.1}, {"capacity": 0}, {"capacity_unit": "unknown"})
        invalid_learning = ({"episodes": 0}, {"alpha": 0}, {"gamma": 1.1}, {"epsilon": -0.1})
        for parameters in invalid_problems:
            with self.subTest(parameters=parameters), self.assertRaises(ValueError):
                ProblemParameters(**parameters)
        for parameters in invalid_learning:
            with self.subTest(parameters=parameters), self.assertRaises(ValueError):
                LearningParameters(**parameters)

    def test_table_and_traceability(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blocks.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "105 Bloques"
            sheet.append(["ID", "X", "Y", "Z", "Valor"])
            sheet.append([8, 1, 2, 3, -4])
            workbook.save(path)
            mine = load_mine(path)
            self.assertEqual(mine.blocks[0].source_row, 2)
            self.assertEqual(mine.blocks[0].id, 8)
            self.assertEqual(len(mine.data_sha256), 64)
            for bad in (None, "4", float("inf"), True):
                sheet["E2"] = bad
                workbook.save(path)
                with self.assertRaises(ValueError):
                    load_mine(path)
            sheet["E2"] = -4
            sheet.append([8, 2, 2, 3, 5])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "duplicados"):
                load_mine(path)

    def test_required_headers_and_coordinates(self):
        with self.assertRaisesRegex(ValueError, "ID interno"):
            Mine((Block(10.0, 1, 1, 1, 3),))
        with self.assertRaisesRegex(ValueError, "duplicados"):
            Mine((Block(1, 1, 1, 1, 3), Block(2, 1, 1, 1, 4)))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.xlsx"
            workbook = Workbook()
            workbook.active.title = "105 Bloques"
            workbook.active.append(["ID", "X", "Y", "Z"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "Encabezados"):
                load_mine(path)

    def test_theme_and_rgb_fill(self):
        from openpyxl.writer.theme import theme_xml
        workbook = Workbook()
        workbook.loaded_theme = theme_xml
        a, b = workbook.active["A1"], workbook.active["B1"]
        a.fill = PatternFill("solid", fgColor=Color(theme=8))
        b.fill = PatternFill("solid", fgColor="004BACC6")
        self.assertEqual(fill_signature(a, workbook), fill_signature(b, workbook))

    def test_tint_uses_luminance_and_indexed_palette(self):
        workbook = Workbook()
        # RGB (0.2, 0.4, 0.6): HLS lightness 0.4 -> 0.7, saturation unchanged.
        self.assertEqual(color_rgb(Color(rgb="00336699", tint=0.5), workbook), "8CB2D9")
        self.assertEqual(color_rgb(Color(rgb="00336699", tint=1), workbook), "FFFFFF")
        self.assertEqual(color_rgb(Color(rgb="00336699", tint=-1), workbook), "000000")
        self.assertEqual(color_rgb(Color(indexed=5), workbook), "FFFF00")


class EvaluationTests(unittest.TestCase):
    def setUp(self):
        self.mine = small_mine()
        self.precedences = grid9_z_up(self.mine)

    def test_same_period_order_and_discount(self):
        plan = Plan((Extraction(10), Extraction(20)))
        result = evaluate_plan(self.mine, plan, ProblemParameters(), self.precedences)
        self.assertTrue(result.feasible)
        self.assertEqual(result.npv, 11)
        delayed = Plan((Extraction(10, 0), Extraction(20, 1)))
        result = evaluate_plan(self.mine, delayed, ProblemParameters(periods=2, discount_rate=0.5), self.precedences)
        self.assertEqual(result.npv, 6)

    def test_errors_no_repair(self):
        for plan in (Plan((Extraction(20),)), Plan((Extraction(20), Extraction(10))),
                     Plan((Extraction(10), Extraction(10))), Plan((Extraction(999),)),
                     Plan((Extraction(10, -1),)), Plan((Extraction(True),))):
            self.assertFalse(evaluate_plan(self.mine, plan, ProblemParameters(), self.precedences).feasible)

    def test_period_order_horizon_and_types(self):
        problem = ProblemParameters(periods=2)
        for plan in (Plan((Extraction(10, 1), Extraction(20, 0))),
                     Plan((Extraction(10, 2),)), Plan((Extraction(10, True),)),
                     Plan((Extraction(10, 0.0),))):
            self.assertFalse(evaluate_plan(self.mine, plan, problem, self.precedences).feasible)

    def test_capacity_units_and_unknown_precedence(self):
        plan = Plan((Extraction(10), Extraction(20)))
        for unit, capacity, feasible in (("blocks", 2, True), ("tonnes", 4, False), ("tonnes", 5, True)):
            result = evaluate_plan(self.mine, plan, ProblemParameters(capacity=capacity, capacity_unit=unit), self.precedences)
            self.assertEqual(result.feasible, feasible)
        self.assertIsNone(evaluate_plan(self.mine, plan, ProblemParameters(), None).feasible)
        with self.assertRaisesRegex(ValueError, "tonelaje"):
            evaluate_plan(Mine((Block(1, 1, 1, 1, 1),)), Plan(),
                          ProblemParameters(capacity=5, capacity_unit="tonnes"), None)

    def test_selection_order_is_explicit_and_checks_closure(self):
        self.assertEqual(order_selection((20, 10), self.precedences), (10, 20))
        with self.assertRaises(ValueError):
            order_selection((20,), self.precedences)

    def test_grid9_and_cycle(self):
        mine = Mine(tuple(Block(i + 1, i % 3, i // 3, 2, 1) for i in range(9)) + (Block(10, 1, 1, 1, 9),))
        self.assertEqual(grid9_z_up(mine).predecessors[10], frozenset(range(1, 10)))
        cycle = Precedences({10: frozenset({20}), 20: frozenset({10})}, "test", "test")
        with self.assertRaisesRegex(ValueError, "Ciclo"):
            cycle.validate(self.mine)


if __name__ == "__main__":
    unittest.main()
