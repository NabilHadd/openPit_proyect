"""Visual metadata and explicit, provisional geometry. Never writes the workbook."""

import colorsys
from dataclasses import dataclass
import json
from pathlib import Path
from types import MappingProxyType
from typing import Mapping
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from src.data.dataset import Mine, file_hash, integer, number


def color_rgb(color, workbook) -> str:
    if color.type == "rgb":
        rgb = color.rgb[-6:]
    elif color.type == "theme":
        if not workbook.loaded_theme:
            raise ValueError("Color de tema sin tema Excel")
        root = ET.fromstring(workbook.loaded_theme)
        scheme = root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme")
        if scheme is None or color.theme >= len(scheme):
            raise ValueError("Color de tema desconocido")
        element = list(scheme[color.theme])[0]
        rgb = element.attrib.get("lastClr", element.attrib.get("val"))
    elif color.type == "indexed" and color.indexed < len(workbook._colors):
        rgb = workbook._colors[color.indexed][-6:]
    else:
        raise ValueError(f"Color no resoluble: {color}")
    channels = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    tint = color.tint
    if tint:
        # SpreadsheetML tint changes HLS luminance, not each RGB channel.
        hue, luminance, saturation = colorsys.rgb_to_hls(*(c / 255 for c in channels))
        luminance = luminance * (1 + tint) if tint < 0 else luminance * (1 - tint) + tint
        channels = [round(c * 255) for c in colorsys.hls_to_rgb(hue, luminance, saturation)]
    return "".join(f"{c:02X}" for c in channels)


def fill_signature(cell, workbook) -> tuple:
    fill = cell.fill
    if fill.patternType != "solid":
        raise ValueError(f"{cell.coordinate}: solo se admiten rellenos sólidos")
    return fill.patternType, color_rgb(fill.fgColor, workbook)


@dataclass(frozen=True)
class Selection:
    sheet: str
    ids: tuple[int, ...]
    cells: tuple[str, ...]
    legend_cell: str = "K3"


@dataclass(frozen=True)
class WorkbookSelections:
    candidates: Selection
    reference: Selection
    declared_value: float
    declared_candidates: int


def read_selection(workbook, sheet_name: str, mine: Mine) -> Selection:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Falta la hoja visual {sheet_name}")
    sheet = workbook[sheet_name]
    legend = fill_signature(sheet["K3"], workbook)
    by_position = {(b.x, b.y, b.z): b.id for b in mine.blocks}
    ids, cells, seen = [], [], set()
    # This adapter recognizes the workbook's labelled level panels, not style IDs.
    panels = []
    for row in sheet:
        for cell in row:
            if cell.column == 1 and isinstance(cell.value, str) and cell.value.startswith("Z="):
                try:
                    z = int(cell.value[2:])
                except ValueError as exc:
                    raise ValueError(f"Nivel no reconocido: {cell.value}") from exc
                header_row = cell.row + 1 if cell.row == 1 else cell.row
                panels.append((z, header_row))
    if not panels:
        raise ValueError(f"{sheet_name}: no se encontraron paneles Z=")
    for z, header_row in panels:
        for row in range(header_row + 1, header_row + 6):
            y = integer(sheet.cell(row, 2).value, f"{sheet_name}:Y fila {row}")
            for column in range(3, 10):
                x = integer(sheet.cell(header_row, column).value, f"{sheet_name}:X columna {column}")
                cell = sheet.cell(row, column)
                block_id = integer(cell.value, f"{sheet_name}!{cell.coordinate}")
                if by_position.get((x, y, z)) != block_id:
                    raise ValueError(f"{sheet_name}!{cell.coordinate}: ID/coordenadas no corresponden a tabla")
                if block_id in seen:
                    raise ValueError(f"Bloque repetido en hoja visual: {block_id}")
                seen.add(block_id)
                if cell.fill.patternType and fill_signature(cell, workbook) == legend:
                    ids.append(block_id)
                    cells.append(cell.coordinate)
    if seen != set(mine.by_id):
        raise ValueError(f"{sheet_name}: la vista no corresponde al yacimiento completo")
    return Selection(sheet_name, tuple(ids), tuple(cells))


def load_selections(path: Path, mine: Mine) -> WorkbookSelections:
    if mine.data_sha256 != file_hash(Path(path)):
        raise ValueError("Las selecciones y el yacimiento deben provenir del mismo archivo")
    workbook = load_workbook(path, data_only=True)
    try:
        return WorkbookSelections(
            read_selection(workbook, "X-Y View", mine),
            read_selection(workbook, "Resultado", mine),
            number(workbook["Resultado"]["L13"].value, "Resultado!L13"),
            integer(workbook["X-Y View"]["K6"].value, "X-Y View!K6"),
        )
    finally:
        workbook.close()


@dataclass(frozen=True)
class Precedences:
    predecessors: Mapping[int, frozenset[int]]
    rule: str
    source: str
    provisional: bool = False

    def __post_init__(self):
        object.__setattr__(self, "predecessors", MappingProxyType({
            k: frozenset(v) for k, v in self.predecessors.items()
        }))

    def validate(self, mine: Mine):
        ids = set(mine.by_id)
        if set(self.predecessors) != ids:
            raise ValueError("Las precedencias deben definir todos los bloques")
        done = set()
        for block_id, predecessors in self.predecessors.items():
            if block_id in predecessors or not predecessors <= ids:
                raise ValueError("Predecesores desconocidos o autorreferencia")
        while len(done) < len(ids):
            ready = {i for i in ids - done if self.predecessors[i] <= done}
            if not ready:
                raise ValueError("Ciclo de precedencias")
            done.update(ready)


def grid9_z_up(mine: Mine) -> Precedences:
    """Higher Z is shallower; 3x3 directly above with provisional edge clipping."""
    return _grid9(mine, 1, "grid9-z-up")


def grid9_z_down(mine: Mine) -> Precedences:
    """Alternative orientation: lower Z is shallower, as numbered in the paper."""
    return _grid9(mine, -1, "grid9-z-down")


def _grid9(mine: Mine, dz: int, rule: str) -> Precedences:
    positions = {(b.x, b.y, b.z): b.id for b in mine.blocks}
    predecessors = {}
    for b in mine.blocks:
        predecessors[b.id] = frozenset(
            positions[(b.x + dx, b.y + dy, b.z + dz)]
            for dx in (-1, 0, 1) for dy in (-1, 0, 1)
            if (b.x + dx, b.y + dy, b.z + dz) in positions
        )
    result = Precedences(predecessors, rule,
                         "Quelopana/Navarra 2021, Fig. 4: nueve superiores; orientación indicada por regla, bordes pendientes", True)
    result.validate(mine)
    return result


def complete_cone_candidates(mine: Mine, surface: str) -> tuple[int, ...]:
    """Diagnostic geometry only: blocks with every overlying cone cell present.

    Unlike clipped-edge precedences, this requires full support inside the grid.
    The result is NOT used to filter the agent or to interpret Excel colours.
    """
    if surface not in {"min-z", "max-z"}:
        raise ValueError("surface debe ser min-z o max-z")
    surface_z = (min if surface == "min-z" else max)(b.z for b in mine.blocks)
    dz = -1 if surface == "min-z" else 1
    positions = {(b.x, b.y, b.z) for b in mine.blocks}
    return tuple(sorted(b.id for b in mine.blocks if all(
        (b.x + dx, b.y + dy, b.z + dz * distance) in positions
        for distance in range(1, abs(surface_z - b.z) + 1)
        for dx in range(-distance, distance + 1) for dy in range(-distance, distance + 1)
    )))


def order_selection(ids, precedences: Precedences) -> tuple[int, ...]:
    """Construct a witness order, NOT the order supplied by a visual selection."""
    remaining, ordered = set(ids), []
    while remaining:
        ready = sorted(i for i in remaining if precedences.predecessors[i] <= set(ordered))
        if not ready:
            raise ValueError("Selección no cerrada bajo precedencias (o ciclo)")
        ordered.extend(ready)
        remaining.difference_update(ready)
    return tuple(ordered)


def load_precedences(path: Path, mine: Mine) -> Precedences:
    """Explicit rules supplied by the caller, not inferred from reference labels."""
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Las precedencias JSON deben ser un objeto")
    if not isinstance(payload.get("source"), str) or not payload["source"].strip():
        raise ValueError("El archivo de precedencias requiere source")
    entries = payload.get("predecessors")
    if not isinstance(entries, dict):
        raise ValueError("predecessors debe ser un objeto ID -> lista de IDs")
    mapping = {}
    for key, values in entries.items():
        if not isinstance(values, list):
            raise ValueError("Cada lista de predecesores debe ser una lista")
        block_id = int(key)
        if str(block_id) != key:
            raise ValueError("Las claves ID deben ser enteros canónicos")
        mapping[block_id] = frozenset(integer(v, "predecessor") for v in values)
    result = Precedences(mapping, "explicit-file", payload["source"])
    result.validate(mine)
    return result
