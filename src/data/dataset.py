"""Read-only ingestion of the block table; no reference selections enter a mine."""

from dataclasses import dataclass
from hashlib import sha256
import math
from numbers import Real
from pathlib import Path
from types import MappingProxyType
from typing import Mapping

from openpyxl import load_workbook


def number(value, field: str) -> float:
    if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value):
        raise ValueError(f"{field}: se requiere un número finito, recibido {value!r}")
    return float(value)


def integer(value, field: str) -> int:
    numeric = number(value, field)
    if not numeric.is_integer():
        raise ValueError(f"{field}: se requiere un entero")
    return int(numeric)


@dataclass(frozen=True)
class Block:
    id: int
    x: int
    y: int
    z: int
    value: float
    source_row: int | None = None
    tonnes: float | None = None


@dataclass(frozen=True)
class Mine:
    blocks: tuple[Block, ...]
    source: str = "synthetic"
    data_sha256: str | None = None
    sheet: str = "105 Bloques"

    def __post_init__(self):
        if not self.blocks:
            raise ValueError("El yacimiento no contiene bloques")
        ids, coordinates = set(), set()
        for block in self.blocks:
            if isinstance(block.id, bool) or not isinstance(block.id, int):
                raise ValueError("ID interno debe ser int; la lectura convierte enteros de Excel")
            position = tuple(integer(getattr(block, c), c) for c in ("x", "y", "z"))
            number(block.value, "Valor")
            if block.id in ids or position in coordinates:
                raise ValueError("ID o coordenadas duplicados")
            if block.tonnes is not None and number(block.tonnes, "tonnes") <= 0:
                raise ValueError("tonnes debe ser positivo")
            ids.add(block.id)
            coordinates.add(position)

    @property
    def by_id(self) -> Mapping[int, Block]:
        return MappingProxyType({block.id: block for block in self.blocks})


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def load_mine(path: Path, sheet: str = "105 Bloques") -> Mine:
    path = Path(path).resolve(strict=True)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Falta la hoja {sheet!r}")
        rows = workbook[sheet].iter_rows(values_only=True)
        headers = next(rows, ())
        required = ("ID", "X", "Y", "Z", "Valor")
        if any(headers.count(name) != 1 for name in required):
            raise ValueError(f"Encabezados obligatorios y únicos: {required}")
        indices = {name: headers.index(name) for name in required}
        blocks = []
        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue
            fields = {name: row[index] for name, index in indices.items()}
            blocks.append(Block(
                id=integer(fields["ID"], f"ID fila {row_number}"),
                x=integer(fields["X"], f"X fila {row_number}"),
                y=integer(fields["Y"], f"Y fila {row_number}"),
                z=integer(fields["Z"], f"Z fila {row_number}"),
                value=number(fields["Valor"], f"Valor fila {row_number}"),
                source_row=row_number,
            ))
        return Mine(tuple(blocks), str(path), file_hash(path), sheet)
    finally:
        workbook.close()
